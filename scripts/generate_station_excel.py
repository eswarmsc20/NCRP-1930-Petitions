import json
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

with open('data/users.json', 'r', encoding='utf-8') as f:
    users = json.load(f)

wb = openpyxl.Workbook()
ws = wb.active
ws.title = "NCRP Station Logins"

# Ensure grid lines are visible
ws.views.sheetView[0].showGridLines = True

# Title block
ws.merge_cells('A1:H1')
title_cell = ws['A1']
title_cell.value = "SRIKAKULAM DISTRICT POLICE - NCRP / 1930 PORTAL"
title_cell.font = Font(name="Calibri", size=15, bold=True, color="FFFFFF")
title_cell.fill = PatternFill(start_color="0B1E3F", end_color="0B1E3F", fill_type="solid")
title_cell.alignment = Alignment(horizontal="center", vertical="center")

ws.merge_cells('A2:H2')
sub_cell = ws['A2']
sub_cell.value = "Confidential Master Credentials Directory (మొత్తం 36 స్టేషన్లు & అడ్మిన్ లాగిన్ వివరాలు)"
sub_cell.font = Font(name="Calibri", size=11, bold=True, color="1E3A8A")
sub_cell.fill = PatternFill(start_color="E0E7FF", end_color="E0E7FF", fill_type="solid")
sub_cell.alignment = Alignment(horizontal="center", vertical="center")

headers = [
    "S.No (క్ర.సం)",
    "Unit / Police Station Name (విభాగం / స్టేషన్ పేరు)",
    "Role (హోదా)",
    "Station Code",
    "Circle (సర్కిల్)",
    "Sub-Division (సబ్ డివిజన్)",
    "User ID (యూజర్ ఐడీ)",
    "Default Password (ప్రారంభ పాస్‌వర్డ్)"
]

ws.append([]) # Blank row 3
ws.append(headers) # Row 4

header_fill = PatternFill(start_color="1E3A8A", end_color="1E3A8A", fill_type="solid")
header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
thin_border = Border(
    left=Side(style='thin', color='D1D5DB'),
    right=Side(style='thin', color='D1D5DB'),
    top=Side(style='thin', color='D1D5DB'),
    bottom=Side(style='thin', color='D1D5DB')
)

for col_num in range(1, len(headers) + 1):
    cell = ws.cell(row=4, column=col_num)
    cell.fill = header_fill
    cell.font = header_font
    cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    cell.border = thin_border

admin_fill = PatternFill(start_color="FEF3C7", end_color="FEF3C7", fill_type="solid")
station_fill_even = PatternFill(start_color="F9FAFB", end_color="F9FAFB", fill_type="solid")
station_fill_odd = PatternFill(start_color="FFFFFF", end_color="FFFFFF", fill_type="solid")

for idx, u in enumerate(users, start=1):
    role_desc = "District Admin" if u.get('role') == 'admin' else "Station User"
    username = u.get('username')
    
    if u.get('role') == 'admin':
        pwd = "Admin@SklmNCRP2026!" if username == "admin" else "Cyber@Sklm2026!"
    else:
        pwd = f"{username}@skl26"
        
    row_data = [
        idx,
        u.get('ps_name'),
        role_desc,
        u.get('station_code'),
        u.get('circle'),
        u.get('sub_division'),
        username,
        pwd
    ]
    ws.append(row_data)
    row_idx = ws.max_row
    
    row_fill = admin_fill if u.get('role') == 'admin' else (station_fill_even if idx % 2 == 0 else station_fill_odd)
    
    for c_idx in range(1, len(row_data) + 1):
        c = ws.cell(row=row_idx, column=c_idx)
        c.fill = row_fill
        c.border = thin_border
        c.font = Font(name="Calibri", size=10, bold=(u.get('role') == 'admin' or c_idx in [1, 7, 8]))
        if c_idx in [1, 3, 4]:
            c.alignment = Alignment(horizontal="center", vertical="center")
        elif c_idx in [7, 8]:
            c.alignment = Alignment(horizontal="center", vertical="center")
        else:
            c.alignment = Alignment(horizontal="left", vertical="center")

# Auto-fit columns
for col in ws.columns:
    max_len = 0
    col_letter = get_column_letter(col[0].column)
    for cell in col:
        if cell.row > 2 and cell.value:
            max_len = max(max_len, len(str(cell.value)))
    ws.column_dimensions[col_letter].width = max(max_len + 4, 12)

ws.column_dimensions['B'].width = 32
ws.column_dimensions['E'].width = 26
ws.column_dimensions['F'].width = 24
ws.column_dimensions['G'].width = 22
ws.column_dimensions['H'].width = 26

output_file = "Srikakulam_Police_NCRP_Station_Logins.xlsx"
wb.save(output_file)
print(f"Successfully generated {output_file} with {len(users)} login records.")
